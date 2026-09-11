import json
from collections import OrderedDict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "上海自考13000英语专升本_高频单词及词组_2026下半年.xlsx"
OUTPUT = ROOT / "assets" / "data" / "vocab_seed.json"


def text(value):
    if value is None:
        return ""
    return str(value).strip()


def normalize(value):
    return " ".join(text(value).lower().split())


def unique_values(values):
    result = []
    for value in values:
        if value and value not in result:
            result.append(value)
    return result


def collect_words(sheet):
    groups = OrderedDict()
    source_rows = 0
    blank_collocations = 0

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not any(value is not None and text(value) for value in row):
            continue

        source_rows += 1
        sequence, term, phonetic, part_of_speech, meaning, collocations, example = row
        term = text(term)
        if not term:
            continue

        key = normalize(term)
        item = groups.setdefault(
            key,
            {
                "id": f"word-{int(sequence):04d}",
                "kind": "word",
                "term": term,
                "normalizedTerm": key,
                "phonetic": text(phonetic),
                "senses": [],
                "sourceRows": [],
            },
        )
        item["phonetic"] = item["phonetic"] or text(phonetic)
        item["sourceRows"].append(int(sequence))
        if not text(collocations):
            blank_collocations += 1
        item["senses"].append(
            {
                "partOfSpeech": text(part_of_speech),
                "meaning": text(meaning),
                "collocations": text(collocations),
                "example": text(example),
            }
        )

    for item in groups.values():
        item["meaning"] = "；".join(
            unique_values(sense["meaning"] for sense in item["senses"])
        )
        item["collocations"] = " / ".join(
            unique_values(
                sense["collocations"]
                for sense in item["senses"]
                if sense["collocations"]
            )
        )
        item["example"] = next(
            (sense["example"] for sense in item["senses"] if sense["example"]),
            "",
        )
        item["partOfSpeech"] = " / ".join(
            unique_values(
                sense["partOfSpeech"]
                for sense in item["senses"]
                if sense["partOfSpeech"]
            )
        )

    return list(groups.values()), source_rows, blank_collocations


def collect_phrases(sheet):
    groups = OrderedDict()
    source_rows = 0

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not any(value is not None and text(value) for value in row):
            continue

        source_rows += 1
        sequence, category, term, meaning, example = row
        term = text(term)
        if not term:
            continue

        key = normalize(term)
        item = groups.setdefault(
            key,
            {
                "id": f"phrase-{int(sequence):04d}",
                "kind": "phrase",
                "term": term,
                "normalizedTerm": key,
                "category": text(category),
                "senses": [],
                "sourceRows": [],
            },
        )
        item["sourceRows"].append(int(sequence))
        item["category"] = item["category"] or text(category)
        item["senses"].append(
            {
                "meaning": text(meaning),
                "example": text(example),
            }
        )

    for item in groups.values():
        item["meaning"] = "；".join(
            unique_values(sense["meaning"] for sense in item["senses"])
        )
        item["example"] = next(
            (sense["example"] for sense in item["senses"] if sense["example"]),
            "",
        )

    return list(groups.values()), source_rows


def collect_exam_info(sheet):
    info = {}
    for row in sheet.iter_rows(min_row=1, values_only=True):
        key = text(row[1]) if len(row) > 1 else ""
        value = text(row[2]) if len(row) > 2 else ""
        if key and value:
            info[key] = value
    return info


def main():
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    words, word_source_count, blank_collocations = collect_words(
        workbook["高频单词(900)"]
    )
    phrases, phrase_source_count = collect_phrases(workbook["高频词组搭配"])
    exam_info = collect_exam_info(workbook["考试信息与备考指南"])

    payload = {
        "version": "2026-second-half",
        "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "sourceFile": SOURCE.name,
        "metadata": {
            "title": "上海自考 13000 英语（专升本）",
            "subtitle": "2026年下半年 · 高频单词及词组",
            "courseCode": exam_info.get("课程代码", "13000"),
            "courseName": exam_info.get("课程名称", "英语（专升本）"),
            "examDateLabel": exam_info.get("考试时间", "2026年10月24日-25日"),
            "examDate": "2026-10-24",
            "wordSourceCount": word_source_count,
            "phraseSourceCount": phrase_source_count,
            "wordCardCount": len(words),
            "phraseCardCount": len(phrases),
            "blankCollocationCount": blank_collocations,
        },
        "items": words + phrases,
    }

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(
        f"Generated {len(words)} word cards and {len(phrases)} phrase cards "
        f"from {word_source_count} + {phrase_source_count} source rows."
    )


if __name__ == "__main__":
    main()
